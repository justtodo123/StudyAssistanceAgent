"""Runtime-generated minimal fixtures for the frozen M7 parser matrix."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


def markdown_bytes() -> bytes:
    return "# 第一章\n\n中文 mixed １２３ 😀\n\n## 简体/繁體\n\n第二段。\n".encode("utf-8")


def text_bytes() -> bytes:
    return "纯文本标题\n\n中文 mixed 123 😀\n".encode("utf-8")


def pdf_bytes() -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, NumberObject, StreamObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_ref = writer._add_object(font)
    resources = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})})
    page[NameObject("/Resources")] = resources
    stream = StreamObject()
    stream._data = b"BT /F1 18 Tf 40 120 Td (PDF page one) Tj ET"
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def pptx_bytes() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Visible slide"
    slide.placeholders[1].text = "Slide body 中文"
    hidden = presentation.slides.add_slide(presentation.slide_layouts[1])
    hidden.shapes.title.text = "Hidden slide"
    hidden.placeholders[1].text = "Must not be indexed"
    hidden._element.set("show", "0")
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("Preamble text")
    document.add_heading("第一节", level=1)
    document.add_paragraph("Section body 中文")
    document.add_heading("小节", level=2)
    document.add_paragraph("Subsection body")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def fixture_bytes(format: str) -> bytes:
    return {
        "md": markdown_bytes,
        "txt": text_bytes,
        "pdf": pdf_bytes,
        "pptx": pptx_bytes,
        "docx": docx_bytes,
    }[format]()


def write_fixture(tmp_path: Path, format: str) -> Path:
    path = tmp_path / f"fixture.{format}"
    path.write_bytes(fixture_bytes(format))
    return path

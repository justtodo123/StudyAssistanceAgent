"""Run reproducible M7 parser/normalization/identity evidence.

The runner creates all input documents under a temporary directory, parses and
normalizes them through the frozen M7 contracts, and compares cold-process and
restart-process evidence. It never discovers or reads an external source tree.

Examples::

    python tools/run_m7_parser_evidence.py --report reports/m7-parser-evidence.json
    python tools/run_m7_parser_evidence.py --seed 20260904 --runs 20
    python tools/run_m7_parser_evidence.py --format pdf --quick

The report is intentionally JSON so it can be archived as evidence without
committing generated fixtures or binary files.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import os
from pathlib import Path
import platform
import subprocess
import sys
import tempfile
import time
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
PLATFORM = ROOT / "platform"
if str(PLATFORM) not in sys.path:
    sys.path.insert(0, str(PLATFORM))

from app.normalized_document import (  # noqa: E402
    CHUNK_SCHEMA_VERSION,
    NORMALIZED_DOCUMENT_SCHEMA,
    normalize_document,
)
from app.parser_matrix import (  # noqa: E402
    PARSER_MATRIX_SCHEMA,
    PARSER_SPECS,
    ParserMatrixError,
    parse_file,
    parser_availability,
)

SCHEMA = "sa.source.parser-normalized-identity-evidence.v1"
DEFAULT_SEED = 20260904
DEFAULT_RUNS = 20
DEFAULT_FIXTURES = 100
DEFAULT_SOURCE_ID = "user-01890f52-47e7-7abc-8def-0123456789ab"
FORMATS = ("md", "txt", "pdf", "pptx", "docx")


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _canonical(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _text(seed: int, index: int, fmt: str) -> str:
    return (
        f"M7 fixture seed {seed} format {fmt} index {index:03d}.\n"
        f"Stable identity content {seed ^ index:08x}; parser evidence only."
    )


def _markdown(seed: int, index: int) -> bytes:
    body = _text(seed, index, "md")
    return f"# Fixture {index:03d}\n\n{body}\n\n## Details\n\nordinal {index}.\n".encode("utf-8")


def _txt(seed: int, index: int) -> bytes:
    return (_text(seed, index, "txt") + "\n").encode("utf-8")


def _pdf(seed: int, index: int) -> bytes:
    # A small deterministic one-page PDF. ASCII content keeps pypdf extraction
    # independent of platform font encodings while still exercising the PDF parser.
    value = (_text(seed, index, "pdf").replace("\n", " ")).encode("ascii")
    stream = b"BT /F1 11 Tf 50 760 Td (" + value.replace(b"(", b"\\(").replace(b")", b"\\)") + b") Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(result))
        result.extend(f"{number} 0 obj\n".encode("ascii"))
        result.extend(obj)
        result.extend(b"\nendobj\n")
    xref = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    result.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    result.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    return bytes(result)


def _pptx(seed: int, index: int) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    from io import BytesIO

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = f"Fixture {index:03d}"
    slide.placeholders[1].text = _text(seed, index, "pptx")
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _docx(seed: int, index: int) -> bytes:
    from docx import Document
    from io import BytesIO

    document = Document()
    document.add_paragraph(f"Preamble {_text(seed, index, 'docx')}")
    document.add_heading(f"Fixture {index:03d}", level=1)
    document.add_paragraph(_text(seed, index, "docx"))
    document.add_heading("Details", level=2)
    document.add_paragraph(f"ordinal {index}.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def fixture_bytes(fmt: str, seed: int, index: int) -> bytes:
    if fmt == "md":
        return _markdown(seed, index)
    if fmt == "txt":
        return _txt(seed, index)
    if fmt == "pdf":
        return _pdf(seed, index)
    if fmt == "pptx":
        return _pptx(seed, index)
    if fmt == "docx":
        return _docx(seed, index)
    raise ValueError(fmt)


def _document_id(source_id: str, logical_uri: str) -> str:
    return _sha256(f"{source_id}\0{logical_uri}".encode("utf-8"))[:32]


def _evidence(path: Path, fmt: str, seed: int, source_id: str) -> dict[str, Any]:
    logical_uri = f"generated/{fmt}/fixture-{path.stem}." + fmt
    data = path.read_bytes()
    parsed = parse_file(path, fmt)
    document = normalize_document(
        parsed,
        source_id=source_id,
        document_id=_document_id(source_id, logical_uri),
        logical_uri=logical_uri,
        format=fmt,
        content_fingerprint=_sha256(data),
        parser_id=parsed.parser_id,
        parser_version=parsed.parser_version,
    )
    canonical = document.canonical_bytes()
    chunks = document.chunks()
    identity = {
        "source_id": source_id,
        "logical_uri": logical_uri,
        "document_id": document.document_id,
        "content_fingerprint": _sha256(data),
        "normalized_text_digest": document.normalized_text_digest,
        "chunk_schema": CHUNK_SCHEMA_VERSION,
        "chunk_ids": [chunk.chunk_id for chunk in chunks],
        "chunk_keys": [chunk.chunk_key for chunk in chunks],
    }
    return {
        "fixture_digest": _sha256(data),
        "parser_identity": parsed.parser_identity,
        "parsed_digest": _sha256(_canonical({
            "format": parsed.format,
            "parser_id": parsed.parser_id,
            "parser_version": parsed.parser_version,
            "units": [
                {
                    "unit_kind": unit.unit_kind,
                    "ordinal": unit.ordinal,
                    "text": unit.text,
                    "title": unit.title,
                    "heading_path": list(unit.heading_path),
                    "visible": unit.visible,
                    "source_ordinal": unit.source_ordinal,
                    "heading_level": unit.heading_level,
                }
                for unit in parsed.units
            ],
        })),
        "normalized_digest": _sha256(canonical),
        "normalized_text_digest": document.normalized_text_digest,
        "document_id": document.document_id,
        "chunk_ids": identity["chunk_ids"],
        "chunk_keys": identity["chunk_keys"],
        "identity_digest": _sha256(_canonical(identity)),
        "unit_count": len(document.units),
        "chunk_count": len(chunks),
    }


def _prepare_fixtures(root: Path, formats: tuple[str, ...], seed: int, count: int) -> dict[str, list[Path]]:
    paths: dict[str, list[Path]] = {}
    for fmt in formats:
        directory = root / fmt
        directory.mkdir(parents=True, exist_ok=True)
        paths[fmt] = []
        for index in range(count):
            path = directory / f"fixture-{index:03d}.{fmt}"
            path.write_bytes(fixture_bytes(fmt, seed, index))
            paths[fmt].append(path)
    return paths


def _run_once(paths: dict[str, list[Path]], formats: tuple[str, ...], seed: int, source_id: str) -> dict[str, list[dict[str, Any]]]:
    return {fmt: [_evidence(path, fmt, seed, source_id) for path in paths[fmt]] for fmt in formats}


def _child(root: Path, formats: tuple[str, ...], seed: int, count: int, source_id: str) -> int:
    paths = {fmt: sorted((root / fmt).glob(f"fixture-*.{fmt}"))[:count] for fmt in formats}
    print(json.dumps(_run_once(paths, formats, seed, source_id), ensure_ascii=False, sort_keys=True, separators=(",", ":")))
    return 0


def _versions(formats: tuple[str, ...]) -> dict[str, str | None]:
    return {fmt: PARSER_SPECS[fmt].parser_version if parser_availability(fmt) else None for fmt in formats}


def _compare(first: dict[str, list[dict[str, Any]]], second: dict[str, list[dict[str, Any]]]) -> tuple[bool, list[str]]:
    failures: list[str] = []
    for fmt in first:
        if first.get(fmt) != second.get(fmt):
            failures.append(f"{fmt}:cold-restart-mismatch")
    return not failures, failures


def run(args: argparse.Namespace) -> dict[str, Any]:
    formats = tuple(args.format) if args.format else FORMATS
    unavailable = [fmt for fmt in formats if not parser_availability(fmt)]
    report: dict[str, Any] = {
        "schema": SCHEMA,
        "status": "PASS" if not unavailable else "UNAVAILABLE",
        "seed": args.seed,
        "fixture_count_per_format": args.count,
        "runs_per_format": args.runs,
        "formats": list(formats),
        "parser_matrix_schema": PARSER_MATRIX_SCHEMA,
        "normalized_document_schema": NORMALIZED_DOCUMENT_SCHEMA,
        "chunk_schema": CHUNK_SCHEMA_VERSION,
        "python": platform.python_version(),
        "platform": platform.platform(),
        "dependency_versions": {fmt: {
            "parser_id": PARSER_SPECS[fmt].parser_id,
            "required": PARSER_SPECS[fmt].parser_version,
            "available": parser_availability(fmt),
        } for fmt in formats},
        "failure_categories": {
            "parser_unavailable": len(unavailable),
            "parse_failed": 0,
            "normalization_failed": 0,
            "cold_restart_mismatch": 0,
            "deterministic_identity_mismatch": 0,
        },
        "formats_result": {},
        "tmp_only": True,
        "external_source_reads": 0,
    }
    if unavailable:
        for fmt in unavailable:
            report["formats_result"][fmt] = {"status": "SKIPPED_UNAVAILABLE", "runs": 0}
        if not args.allow_unavailable:
            return report

    available_formats = tuple(fmt for fmt in formats if fmt not in unavailable)
    with tempfile.TemporaryDirectory(prefix="m7-parser-evidence-") as temporary:
        root = Path(temporary)
        paths = _prepare_fixtures(root, available_formats, args.seed, args.count)
        for fmt in available_formats:
            result = {"status": "PASS", "runs": 0, "fixture_count": args.count, "fixture_digests": [], "run_digests": []}
            baseline: dict[str, Any] | None = None
            for run_no in range(args.runs):
                started = time.perf_counter()
                try:
                    cold = _run_once(paths, (fmt,), args.seed, DEFAULT_SOURCE_ID)
                    env = os.environ.copy()
                    env["PYTHONPATH"] = str(ROOT) + os.pathsep + str(PLATFORM) + os.pathsep + env.get("PYTHONPATH", "")
                    command = [sys.executable, "-m", "tools.run_m7_parser_evidence", "--child", "--root", str(root), "--format", fmt, "--seed", str(args.seed), "--count", str(args.count)]
                    completed = subprocess.run(command, cwd=ROOT, env=env, check=True, capture_output=True, text=True)
                    restart = json.loads(completed.stdout)
                    same, failures = _compare(cold, restart)
                    if not same:
                        result["status"] = "FAIL"
                        report["failure_categories"]["cold_restart_mismatch"] += len(failures)
                    if baseline is None:
                        baseline = cold
                        result["fixture_digests"] = [item["fixture_digest"] for item in cold[fmt]]
                    elif cold != baseline:
                        result["status"] = "FAIL"
                        report["failure_categories"]["deterministic_identity_mismatch"] += 1
                    result["run_digests"].append(_sha256(_canonical(cold[fmt])))
                    result["runs"] += 1
                except ParserMatrixError:
                    result["status"] = "FAIL"
                    report["failure_categories"]["parse_failed"] += 1
                except Exception:
                    result["status"] = "FAIL"
                    report["failure_categories"]["normalization_failed"] += 1
                result.setdefault("elapsed_ms", []).append(round((time.perf_counter() - started) * 1000, 3))
            report["formats_result"][fmt] = result
    if any(item.get("status") == "FAIL" for item in report["formats_result"].values()):
        report["status"] = "FAIL"
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, help="write the JSON report to this path")
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    parser.add_argument("--runs", type=int, default=DEFAULT_RUNS)
    parser.add_argument("--count", type=int, default=DEFAULT_FIXTURES)
    parser.add_argument("--format", dest="format", choices=FORMATS, action="append", help="limit formats; repeatable")
    parser.add_argument("--quick", action="store_true", help="run one fixture and one run per format")
    parser.add_argument("--allow-unavailable", action="store_true", help="record unavailable formats and continue with available formats")
    parser.add_argument("--child", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--root", type=Path, help=argparse.SUPPRESS)
    args = parser.parse_args(argv)
    if args.quick:
        args.runs, args.count = 1, 1
    if args.runs <= 0 or args.count <= 0:
        parser.error("--runs and --count must be positive")
    formats = tuple(args.format) if args.format else FORMATS
    if args.child:
        if args.root is None:
            parser.error("--child requires --root")
        return _child(args.root, formats, args.seed, args.count, DEFAULT_SOURCE_ID)
    report = run(args)
    payload = json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(payload, encoding="utf-8")
    print(payload, end="")
    return 0 if report["status"] == "PASS" or args.allow_unavailable and report["status"] == "UNAVAILABLE" else 2


if __name__ == "__main__":
    raise SystemExit(main())
